#!/usr/bin/env python3
"""Regenerate the Word template the skill builds on, from assets/brand.json and assets/logo.png.

    python tools/make_template.py

Run it after changing brand.json or the logo, then run tools/package_skill.py.
"""

import json
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parent.parent
SKILL = REPO / "plugins" / "panasa-docs" / "skills" / "project-documentation"
ASSETS = SKILL / "assets"
sys.path.insert(0, str(SKILL / "scripts"))

from build_docx import PPR_ORDER, insert_ordered, make_el  # noqa: E402

from docx import Document  # noqa: E402
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

TEXT = "262626"
GREY = "7F7F7F"
RULE = "BFBFBF"


def set_fonts(rpr, name):
    """Point every script slot at NAME and drop theme-font links, which would override it."""
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in list(fonts.attrib):
        if attr.lower().endswith("theme"):
            del fonts.attrib[attr]
    for slot in ("ascii", "hAnsi", "cs", "eastAsia"):
        fonts.set(qn(f"w:{slot}"), name)


def get_style(doc, name):
    try:
        return doc.styles[name]
    except KeyError:
        style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Normal"]
        return style


def fmt(style, font=None, size=None, bold=None, italic=None, color=None, before=None, after=None,
        line=None, left=None, right=None, first=None, keep_next=None, align=None):
    if font:
        set_fonts(style.element.get_or_add_rPr(), font)
    f = style.font
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color:
        f.color.rgb = RGBColor.from_string(color)
        el = style.element.rPr.find(qn("w:color"))
        for attr in ("themeColor", "themeShade", "themeTint"):
            el.attrib.pop(qn(f"w:{attr}"), None)
    pf = style.paragraph_format
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line
    if left is not None:
        pf.left_indent = Cm(left)
    if right is not None:
        pf.right_indent = Cm(right)
    if first is not None:
        pf.first_line_indent = Cm(first)
    if keep_next is not None:
        pf.keep_with_next = keep_next
    if align is not None:
        pf.alignment = align
    return style


def drop(style, tag):
    ppr = style.element.pPr
    if ppr is not None:
        for el in ppr.findall(qn(tag)):
            ppr.remove(el)


def border(ppr, side, color, size):
    bdr = ppr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = insert_ordered(ppr, OxmlElement("w:pBdr"), PPR_ORDER)
    bdr.append(make_el(f"w:{side}", val="single", sz=size, space=4, color=color))


def fill(ppr, color):
    insert_ordered(ppr, make_el("w:shd", val="clear", color="auto", fill=color), PPR_ORDER)


def add_field(paragraph, instr):
    field = make_el("w:fldSimple", instr=instr)
    run, text = OxmlElement("w:r"), OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def main():
    brand = json.loads((ASSETS / "brand.json").read_text(encoding="utf-8"))
    font = brand["font"]
    heading = brand.get("heading_font") or font
    primary, accent, light = brand["primary_color"], brand["accent_color"], brand["light_color"]
    doc = Document()

    defaults = doc.styles.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    set_fonts(defaults, font)
    lang = defaults.find(qn("w:lang"))
    if lang is not None:
        lang.set(qn("w:val"), brand.get("language", "en-IN"))

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin, sec.bottom_margin = Cm(2.5), Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.header_distance, sec.footer_distance = Cm(1.2), Cm(1.0)
    width = sec.page_width - sec.left_margin - sec.right_margin

    # Text styles
    fmt(get_style(doc, "Normal"), font, 10.5, color=TEXT, after=6, line=1.15)
    drop(fmt(get_style(doc, "Title"), heading, 28, bold=True, color=primary, before=110, after=8), "w:pBdr")
    fmt(get_style(doc, "Subtitle"), heading, 15, italic=False, color=accent, after=30)
    for name, size, color, before, after in (("Heading 1", 16, primary, 18, 6), ("Heading 2", 13, primary, 14, 4),
                                             ("Heading 3", 11.5, accent, 10, 4), ("Heading 4", 10.5, TEXT, 8, 2)):
        fmt(get_style(doc, name), heading, size, bold=True, italic=False, color=color,
            before=before, after=after, keep_next=True)
    fmt(get_style(doc, "Front Heading"), heading, 16, bold=True, color=primary, after=12, keep_next=True)
    fmt(get_style(doc, "Cover Brand"), heading, 24, bold=True, color=primary, before=12, after=0)
    fmt(get_style(doc, "Cover Notice"), size=8.5, italic=True, color=GREY, before=36)
    fmt(get_style(doc, "Table Text"), size=9.5, before=2, after=2, line=1.0)
    fmt(get_style(doc, "Spacer"), size=4, before=0, after=0, line=1.0)
    fmt(get_style(doc, "Figure"), before=6, after=2, keep_next=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    fmt(get_style(doc, "Caption"), size=9, bold=False, italic=True, color=GREY, after=10,
        align=WD_ALIGN_PARAGRAPH.CENTER)

    code = fmt(get_style(doc, "Code Block"), "Consolas", 9, color="1F1F1F", before=2, after=8, line=1.0,
               left=0.2, right=0.2)
    fill(code.element.get_or_add_pPr(), "F2F2F2")
    note = fmt(get_style(doc, "Note"), before=4, after=8, left=0.3, right=0.2)
    fill(note.element.get_or_add_pPr(), light)
    border(note.element.get_or_add_pPr(), "left", accent, 18)

    # Lists: bullets use Word's own list styles; numbers are typed, so each list restarts cleanly.
    for name in ("List Bullet", "List Bullet 2", "List Bullet 3"):
        fmt(get_style(doc, name), after=3)
    for level, name in enumerate(("Numbered Item", "Numbered Item 2", "Numbered Item 3"), 1):
        style = fmt(get_style(doc, name), after=3, left=0.75 * level, first=-0.6)
        style.paragraph_format.tab_stops.add_tab_stop(Cm(0.75 * level))
    check = fmt(get_style(doc, "Checklist"), after=3, left=0.75, first=-0.6)
    check.paragraph_format.tab_stops.add_tab_stop(Cm(0.75))

    # Word's built-in table of contents styles are named "toc 1", "toc 2".
    for level, name in enumerate(("toc 1", "toc 2")):
        style = fmt(get_style(doc, name), bold=level == 0, before=6 if level == 0 else 0, after=2, left=0.5 * level)
        style.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    for name in ("Header", "Footer"):
        drop(fmt(get_style(doc, name), size=8.5, color=GREY, after=0), "w:tabs")

    # Header and footer; the cover page has none. {{...}} tokens are filled in by build_docx.py.
    sec.different_first_page_header_footer = True
    sec.first_page_header.is_linked_to_previous = False
    sec.first_page_footer.is_linked_to_previous = False

    head = sec.header.paragraphs[0]
    head.style = get_style(doc, "Header")
    head.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    logo = ASSETS / "logo.png"
    if logo.is_file():
        head.add_run().add_picture(str(logo), height=Cm(0.8))
    else:
        run = head.add_run(brand["company"].upper())
        run.bold = True
        run.font.color.rgb = RGBColor.from_string(primary)
    head.add_run("\t{{DOC_REF}}")
    border(head._p.get_or_add_pPr(), "bottom", RULE, 4)

    foot = sec.footer.paragraphs[0]
    foot.style = get_style(doc, "Footer")
    foot.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    foot.add_run(brand["company"] + "  ·  {{CLASSIFICATION}}\tPage ")
    add_field(foot, "PAGE")
    foot.add_run(" of ")
    add_field(foot, "NUMPAGES")
    border(foot._p.get_or_add_pPr(), "top", RULE, 4)

    props = doc.core_properties
    props.title = f"{brand['company']} document template"
    props.author = props.last_modified_by = brand["company"]
    props.comments = "Generated by tools/make_template.py"
    out = ASSETS / "panasa-template.docx"
    doc.save(str(out))
    print(f"Wrote {out}" + ("" if logo.is_file() else " (no logo.png found, so the company name is used instead)"))


if __name__ == "__main__":
    main()
