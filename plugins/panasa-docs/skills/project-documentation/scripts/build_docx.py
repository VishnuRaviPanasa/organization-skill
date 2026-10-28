#!/usr/bin/env python3
"""Build a Panasa-branded Word document from a project-documentation Markdown source.

Usage:
    python build_docx.py SOURCE.md [--out FOLDER_OR_FILE] [--check] [--blank] [--force]

    --check   validate only; nothing is written
    --blank   build a fill-in-by-hand template (placeholders and missing details allowed)
    --force   overwrite an existing approved .docx (never needed for drafts)

Exit code 0: valid (and built, unless --check); warnings may still be printed.
Exit code 1: errors found; nothing is written.
Needs python-docx: pip install python-docx
"""

import argparse
import json
import re
import sys
from datetime import datetime, timezone
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX, WD_TAB_ALIGNMENT
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Emu, RGBColor
except ImportError:
    sys.exit("python-docx is not installed. Run: pip install python-docx")

SKILL_DIR = Path(__file__).resolve().parent.parent
ASSETS = SKILL_DIR / "assets"
TEMPLATES = SKILL_DIR / "templates"

DOC_TYPES = {
    "OVR": ("Project Overview", "overview.md"),
    "SRS": ("Software Requirements Specification", "srs.md"),
    "TDD": ("Technical Design Document", "technical-design.md"),
    "RHN": ("Release and Handover Note", "release-handover.md"),
}
STATUSES = ("Draft", "In review", "Approved")
REQUIRED = ("type", "project", "project_code", "client", "version", "status", "date", "prepared_by")
COVER_FIELDS = (
    ("Project", "project"),
    ("Project code", "project_code"),
    ("Project code status", "project_code_status"),
    ("Client", "client"),
    ("Document type", "type_label"),
    ("Version", "version"),
    ("Status", "status"),
    ("Date", "date"),
    ("Prepared by", "prepared_by"),
    ("Reviewed by", "reviewed_by"),
    ("Approved by", "approved_by"),
    ("Approval reference", "approval_ref"),
    ("Classification", "classification"),
)
IMAGE_TYPES = {".png", ".jpg", ".jpeg", ".gif", ".bmp"}
LOGO_TYPES = {".png", ".jpg", ".jpeg"}
CODE_STATUSES = ("agreed", "provisional")
# File names that usually hold credentials; reported at hand-over when found near the document, never opened.
SECRET_FILES = re.compile(r"(?i)^(\.env(\..*)?|.*\.(pem|key|p12|pfx|jks|keystore)|google-services\.json|"
                          r"GoogleService-Info\.plist|credentials.*\.json|id_rsa.*|id_ed25519.*)$")
SKIP_DIRS = {".git", "node_modules", "__pycache__", ".venv", "venv", "dist", "build"}
# Values that don't name a real person or record.
NOT_A_NAME = {"pending", "tbc", "tba", "n/a", "na", "-", "none", "client", "panasa", "yes", "approved", "signed", "unknown"}
STANDARD_IDS = ("FR", "NFR", "OQ", "DEC", "RSK", "KI")

# Word can refuse files whose property elements are out of schema order, so new
# children are inserted at their schema position instead of being appended.
PPR_ORDER = [
    "w:pStyle", "w:keepNext", "w:keepLines", "w:pageBreakBefore", "w:framePr", "w:widowControl",
    "w:numPr", "w:suppressLineNumbers", "w:pBdr", "w:shd", "w:tabs", "w:suppressAutoHyphens",
    "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
    "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection",
    "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr",
    "w:sectPr", "w:pPrChange",
]
TCPR_ORDER = [
    "w:cnfStyle", "w:tcW", "w:gridSpan", "w:hMerge", "w:vMerge", "w:tcBorders", "w:shd",
    "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark",
]
TBLPR_ORDER = [
    "w:tblStyle", "w:tblpPr", "w:tblOverlap", "w:bidiVisual", "w:tblStyleRowBandSize",
    "w:tblStyleColBandSize", "w:tblW", "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders",
    "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription",
]
SETTINGS_ORDER = [
    "w:updateFields", "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat",
    "w:docVars", "w:rsids", "w:attachedSchema", "w:themeFontLang", "w:clrSchemeMapping",
    "w:doNotIncludeSubdocsInStats", "w:doNotAutoCompressPictures", "w:forceUpgrade", "w:captions",
    "w:readModeInkLockDown", "w:smartTagType", "w:shapeDefaults", "w:doNotEmbedSmartTags",
    "w:decimalSymbol", "w:listSeparator",
]

HEADING = re.compile(r"^(#{1,4})\s+(.+?)\s*$")
CHECK = re.compile(r"^(\s*)[-*+]\s+\[([ xX])\]\s+(.*)$")
BULLET = re.compile(r"^(\s*)[-*+]\s+(.*)$")
NUMBER = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
IMAGE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)$")
TABLE_SEP = re.compile(r"^\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?$")
INLINE = re.compile(
    r"\*\*(?P<b>.+?)\*\*"
    r"|`(?P<c>[^`]+)`"
    r"|\[(?P<lt>[^\]]+)\]\((?P<lu>[^)\s]+)\)"
    r"|(?<![\w*])\*(?P<i>[^*\s](?:[^*]*?[^*\s])?)\*(?![\w*])"
)
BR = re.compile(r"<br\s*/?>", re.I)
TBD = re.compile(r"(\bTBD\b(?:\s*\([^)]*\))?)")
PROPOSED = re.compile(r"(\(Proposed\b[^)]*\))")
MARKERS = re.compile(r"(\bTBD\b(?:\s*\([^)]*\))?|\(Proposed\b[^)]*\))")
TBD_OWNER = re.compile(r"\bTBD\s*\(owner:\s*([^)]*?[^)\s][^)]*)\)")
TBD_PHRASE = re.compile(r"\b(?:tbd|Tbd)\b|(?i:\bto be (?:decided|confirmed|determined|defined|agreed)\b)")
# A placeholder starts with a letter: "<500 ms" and "<ops@x.com>" are not placeholders.
PLACEHOLDER = re.compile(r"<(?!br\b)[A-Za-z][^<>\n@]{0,78}>")
OTHER_PLACEHOLDERS = [
    (re.compile(r"\{\{[^{}\n]+\}\}"), "a template token {{...}}"),
    (re.compile(r"\bX{3,}\b"), "an 'XXX' marker"),
    (re.compile(r"\b(?:TODO|FIXME|TBC|TBA)\b"), "a TODO/TBC marker (use TBD (owner: ...))"),
    (re.compile(r"(?i)\blorem ipsum\b"), "filler text (lorem ipsum)"),
]
BRACKET_PLACEHOLDER = re.compile(r"(?<!!)\[(?![ xX]\])([A-Z][a-z]+(?: [A-Za-z]+){0,4})\](?!\()")
NOT_APPLICABLE = re.compile(r"(?i)^(?:n/?a|not applicable)\b\s*(.*)$")
UNDECIDED = re.compile(r"(?i)\b(not (?:yet )?(?:been )?decided|yet to|will be|once |later|pending|at this stage|not (?:yet )?known)\b")
VAGUE = re.compile(r"(?i)\b(fast|quick(?:ly)?|secure|user[- ]friendly|intuitive|easy|simple|robust|seamless|scalable|"
                   r"efficient|flexible|modern|always available|as soon as possible)\b")
DATE = re.compile(r"\d{1,2} (Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec) \d{4}")
ID_DEF = re.compile(r"^((?:%s)-\d{3})\b" % "|".join(STANDARD_IDS))
FR_HEADING = re.compile(r"^(FR-\d{3})\b")
SECRETS = [
    (re.compile(r"-----BEGIN [A-Z ]*PRIVATE KEY-----"), "a private key"),
    (re.compile(r"\b(?:AKIA|ASIA)[0-9A-Z]{16}\b"), "an AWS access key"),
    (re.compile(r"\bAIza[0-9A-Za-z_-]{35}\b"), "a Google API key"),
    (re.compile(r"\bGOCSPX-[A-Za-z0-9_-]{20,}"), "a Google OAuth client secret"),
    # Real sk- keys mix digits and capitals; branch names like sk-offline-sync don't.
    (re.compile(r"\bsk-(?=[A-Za-z0-9_-]*\d)(?=[A-Za-z0-9_-]*[A-Z])[A-Za-z0-9_-]{20,}"), "a secret API key"),
    (re.compile(r"\b[sr]k_(?:live|test)_[0-9A-Za-z]{16,}"), "a Stripe key"),
    (re.compile(r"\b(?:gh[pousr]_[A-Za-z0-9]{30,}|github_pat_[A-Za-z0-9_]{30,})"), "a GitHub token"),
    (re.compile(r"\bglpat-[A-Za-z0-9_-]{20,}"), "a GitLab token"),
    (re.compile(r"\bxox[abprs]-[A-Za-z0-9-]{10,}"), "a Slack token"),
    (re.compile(r"https://hooks\.slack\.com/services/[A-Za-z0-9/]+"), "a Slack webhook URL"),
    (re.compile(r"\bSG\.[A-Za-z0-9_-]{16,}\.[A-Za-z0-9_-]{16,}"), "a SendGrid key"),
    (re.compile(r"\beyJ[A-Za-z0-9_-]{8,}\.eyJ[A-Za-z0-9_-]{8,}\.[A-Za-z0-9_-]{8,}"), "a JSON Web Token"),
    (re.compile(r"(?i)\bAccountKey=[A-Za-z0-9+/=]{20,}"), "an Azure storage key"),
]
URI_PASSWORD = re.compile(r"\b[a-z][a-z0-9+.-]*://[^\s:/@]+:([^\s@/]+)@", re.I)
SECRET_ASSIGNMENT = re.compile(
    r"(?i)(pass(?:word|wd)?|pwd|secret|token|api[_-]?key|access[_-]?key|client[_-]?secret|account[_-]?key|private[_-]?key)"
    r"[A-Za-z0-9_]*[\"']?\s*[:=]\s*[\"']?([^\s\"',;]+)")
SAFE_VALUES = {"stored", "kept", "held", "in", "from", "via", "see", "none", "not", "set", "managed", "provided",
               "required", "redacted", "password", "pass", "pwd", "secret", "xxx", "****"}


# ---------------------------------------------------------------- parsing

def plain(text):
    """TEXT without Markdown inline markers, for widths and comparisons."""
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return BR.sub(" ", text).replace("**", "").replace("`", "").replace("*", "")


def strip_comments(text):
    """Remove <!-- --> comments but keep line numbers stable."""
    return re.sub(r"<!--.*?-->", lambda m: "\n" * m.group(0).count("\n"), text, flags=re.S)


def parse_source(text):
    """Split TEXT into (meta, revisions, body_lines, first_body_line_number, errors)."""
    lines = text.split("\n")
    if not lines or lines[0].strip() != "---":
        return {}, [], lines, 1, ["Line 1: the file must start with front matter between '---' lines."]
    end = next((i for i in range(1, len(lines)) if lines[i].strip() == "---"), None)
    if end is None:
        return {}, [], lines, 1, ["The front matter has no closing '---' line."]
    meta, revisions, errors = {}, [], []
    for i in range(1, end):
        s = lines[i].strip()
        if not s or s.startswith("#"):
            continue
        if ":" not in s:
            errors.append(f"Line {i + 1}: front matter lines must look like 'key: value'.")
            continue
        key, value = (part.strip() for part in s.split(":", 1))
        if len(value) >= 2 and value[0] == value[-1] and value[0] in "\"'":
            value = value[1:-1]
        key = key.lower().replace(" ", "_")
        if key == "revision":
            revisions.append(([f.strip() for f in value.split("|")], i + 1))
        else:
            meta[key] = value
    return meta, revisions, lines[end + 1:], end + 2, errors


def split_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|") and not s.endswith("\\|"):
        s = s[:-1]
    return [cell.strip().replace("\\|", "|") for cell in re.split(r"(?<!\\)\|", s)]


def indent_level(spaces):
    return min(len(spaces.replace("\t", "    ")) // 2, 2)


def list_item(line, ln):
    m = CHECK.match(line)
    if m:
        return {"kind": "check", "done": m.group(2).lower() == "x", "text": m.group(3), "line": ln}
    m = NUMBER.match(line)
    if m:
        return {"kind": "number", "level": indent_level(m.group(1)), "num": m.group(2), "text": m.group(3), "line": ln}
    m = BULLET.match(line)
    return {"kind": "bullet", "level": indent_level(m.group(1)), "text": m.group(2), "line": ln}


def starts_block(line):
    s = line.strip()
    return bool(not s or s.startswith(("```", "|", ">")) or s == "\\newpage" or HEADING.match(line)
                or BULLET.match(line) or NUMBER.match(line) or IMAGE.match(s))


def parse_blocks(lines, first_line):
    """Turn Markdown LINES into block dicts. Line numbers refer to the source file."""
    blocks, i, n = [], 0, len(lines)
    while i < n:
        line = lines[i]
        s, ln = line.strip(), first_line + i
        if not s or s in ("---", "***", "___"):
            i += 1
        elif s.startswith("```"):
            code, i = [], i + 1
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i].rstrip())
                i += 1
            blocks.append({"kind": "code", "lines": code, "line": ln})
            i += 1
        elif HEADING.match(line):
            m = HEADING.match(line)
            blocks.append({"kind": "heading", "level": len(m.group(1)), "text": m.group(2), "line": ln})
            i += 1
        elif s == "\\newpage":
            blocks.append({"kind": "pagebreak", "line": ln})
            i += 1
        elif s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                if not TABLE_SEP.match(lines[i].strip()):
                    rows.append(split_row(lines[i]))
                i += 1
            blocks.append({"kind": "table", "rows": rows, "line": ln})
        elif s.startswith(">"):
            parts = []
            while i < n and lines[i].strip().startswith(">"):
                parts.append(lines[i].strip()[1:].strip())
                i += 1
            blocks.append({"kind": "note", "text": " ".join(p for p in parts if p), "line": ln})
        elif IMAGE.match(s):
            m = IMAGE.match(s)
            blocks.append({"kind": "image", "alt": m.group(1).strip(), "path": m.group(2).strip(), "line": ln})
            i += 1
        elif BULLET.match(line) or NUMBER.match(line):
            block = list_item(line, ln)
            i += 1
            while i < n and lines[i].strip() and lines[i][:1].isspace() and not starts_block(lines[i]):
                block["text"] += " " + lines[i].strip()
                i += 1
            blocks.append(block)
        else:
            parts = [s]
            i += 1
            while i < n and not starts_block(lines[i]):
                parts.append(lines[i].strip())
                i += 1
            blocks.append({"kind": "para", "text": " ".join(parts), "line": ln})
    return blocks


# ---------------------------------------------------------------- checks

def expected_sections(doc_type):
    text = (TEMPLATES / DOC_TYPES[doc_type][1]).read_text(encoding="utf-8")
    _, _, lines, first, _ = parse_source(text)
    blocks = parse_blocks(strip_comments("\n".join(lines)).split("\n"), first)
    return [b["text"] for b in blocks if b["kind"] == "heading" and b["level"] == 1]


def check_structure(blocks, expected, errors, warnings):
    def norm(t):
        return " ".join(plain(t).split()).lower()

    want = [norm(t) for t in expected]
    seen, order, appendix_seen = set(), [], False
    for b in blocks:
        if b["kind"] != "heading" or b["level"] != 1:
            continue
        key = norm(b["text"])
        if key.startswith("appendix"):
            appendix_seen = True
        elif key not in want:
            errors.append(f"Line {b['line']}: '# {b['text']}' is not a standard section. Use the template's "
                          "section titles, put extra material in a '##' sub-section, or add it at the end "
                          "as '# Appendix A: <title>'.")
        elif key in seen:
            errors.append(f"Line {b['line']}: section '{b['text']}' appears twice.")
        else:
            if appendix_seen:
                errors.append(f"Line {b['line']}: '{b['text']}' comes after an appendix. Appendices go at the end.")
            seen.add(key)
            order.append(key)
    for title, key in zip(expected, want):
        if key not in seen:
            errors.append(f"Missing section '# {title}'. Keep every standard section; if one doesn't apply, "
                          "say why in one sentence.")
    if order != [k for k in want if k in seen]:
        errors.append("Sections are out of order. The standard order is: " + "; ".join(expected) + ".")
    if blocks and (blocks[0]["kind"] != "heading" or blocks[0]["level"] != 1):
        warnings.append(f"Line {blocks[0]['line']}: content before the first section doesn't fit the "
                        "standard layout; move it into a section.")
    return len(seen), len(want)


def check_date(value, where, errors, warnings):
    if not DATE.fullmatch(value):
        warnings.append(f"{where}: dates should be written like 27 Oct 2026 (found '{value}').")
        return
    try:
        datetime.strptime(value, "%d %b %Y")
    except ValueError:
        errors.append(f"{where}: '{value}' is not a real date.")


def check_meta(meta, revisions, src_dir, errors, warnings):
    for key in REQUIRED:
        if not meta.get(key):
            errors.append(f"Front matter: '{key}' is missing.")
    for key, value in meta.items():
        if PLACEHOLDER.search(value):
            errors.append(f"Front matter: replace the placeholder in '{key}' ({value}).")
    code, version, status = meta.get("project_code", ""), meta.get("version", ""), meta.get("status", "")
    if code and not re.fullmatch(r"[A-Z][A-Z0-9]{1,11}", code):
        errors.append("project_code must be 2 to 12 capital letters or digits, starting with a letter (e.g. NIRVANA).")
    code_status = meta.get("project_code_status", "").lower()
    if code_status and code_status not in CODE_STATUSES:
        errors.append(f"project_code_status must be 'agreed' or 'provisional' (found '{meta['project_code_status']}').")
    if code_status == "provisional" and status == "Approved":
        errors.append("An approved document needs an agreed project code; it is still marked provisional.")
    if version and not re.fullmatch(r"\d+\.\d+", version):
        errors.append(f"version must look like 0.1 or 1.0 (found '{version}').")
    if status and status not in STATUSES:
        errors.append(f"status must be one of: {', '.join(STATUSES)} (found '{status}').")
    if meta.get("date"):
        check_date(meta["date"], "Front matter date", errors, warnings)
    if status == "Approved" and re.fullmatch(r"\d+\.\d+", version) and int(version.split(".")[0]) < 1:
        errors.append("An approved document must be version 1.0 or higher.")
    logo = meta.get("client_logo")
    if logo:
        path = Path(logo)
        path = path if path.is_absolute() else src_dir / path
        if not path.is_file():
            errors.append(f"Front matter: client_logo file not found: {logo}")
        elif path.suffix.lower() not in LOGO_TYPES:
            errors.append("Front matter: client_logo must be a PNG or JPG file.")
    if not revisions:
        errors.append("Add at least one 'revision: version | date | author | changes' line to the front matter.")
    for fields, ln in revisions:
        if len(fields) != 4 or not all(fields):
            errors.append(f"Line {ln}: a revision line needs four parts: version | date | author | changes.")
        elif any(PLACEHOLDER.search(f) for f in fields):
            errors.append(f"Line {ln}: replace the placeholders in this revision line.")
        else:
            check_date(fields[1], f"Line {ln}", errors, warnings)
    if revisions and version and revisions[-1][0][0] != version:
        warnings.append(f"version is {version} but the last revision line is for {revisions[-1][0][0]}. "
                        f"Add a revision line for {version}.")


def numbers(lines, limit=12):
    return ", ".join(map(str, lines[:limit])) + (", ..." if len(lines) > limit else "")


def real_secret(value):
    """True when VALUE looks like an actual password or key, not a description, variable or example."""
    v = value.strip().strip(".,;)")
    if not v or v.lower() in SAFE_VALUES or v[0] in "<{$%*[(" or v.upper().startswith("TBD"):
        return False
    if len(v) < 6 or re.fullmatch(r"[\d.]+[a-z]{0,3}", v):
        return False
    mixed = re.search(r"[^A-Za-z0-9]", v) and re.search(r"[A-Z]", v) and re.search(r"[a-z]", v)
    return bool(re.search(r"\d", v) or mixed)


def check_secrets(raw_lines, errors):
    """Scan every line of the source, including front matter, comments and code blocks."""
    for ln, line in enumerate(raw_lines, 1):
        found = [label for pattern, label in SECRETS if pattern.search(line)]
        if any(real_secret(m.group(1)) for m in URI_PASSWORD.finditer(line)):
            found.append("a URL with a password in it")
        if any(real_secret(m.group(2)) for m in SECRET_ASSIGNMENT.finditer(line)):
            found.append("a password or key value")
        if found:
            errors.append(f"Line {ln}: looks like {found[0]}. Remove it and say where it is stored instead, "
                          "and tell the user it may need to be rotated.")


def unowned_table_tbds(blocks):
    """Line numbers of table rows with a bare TBD and no owner anywhere in the row.
    A row is owned by an Owner column value or by a TBD (owner: ...) in another cell."""
    found = []
    for b in blocks:
        if b["kind"] != "table" or not b["rows"]:
            continue
        header = [plain(c).strip().lower() for c in b["rows"][0]]
        owner_col = next((k for k, h in enumerate(header) if "owner" in h), None)
        for i, row in enumerate(b["rows"][1:], 1):
            cells = [plain(c) for c in row]
            bare = sum(len(re.findall(r"\bTBD\b", c)) - len(TBD_OWNER.findall(c)) for c in cells)
            owner = cells[owner_col].strip() if owner_col is not None and owner_col < len(cells) else ""
            if bare > 0 and not any(TBD_OWNER.search(c) for c in cells) and not (owner and "TBD" not in owner):
                found.append(b["line"] + i + 1)  # header, separator, then rows
    return found


def check_body(blocks, lines, first_line, src_dir, blank, errors, warnings, tbds, proposed):
    placeholders, no_owner, phrases, brackets, in_code = [], [], [], [], False
    for idx, line in enumerate(lines):
        ln = first_line + idx
        if line.strip().startswith("```"):
            in_code = not in_code
            continue
        if in_code:
            continue
        text = re.sub(r"`[^`]*`", "", line)  # inline code is literal, never a placeholder or marker
        if TBD.search(text):
            tbds.append((ln, line.strip()))
        if PROPOSED.search(text):
            proposed.append((ln, line.strip()))
        if blank:
            continue
        if PLACEHOLDER.search(text):
            placeholders.append(ln)
        for pattern, label in OTHER_PLACEHOLDERS:
            if pattern.search(text):
                errors.append(f"Line {ln}: contains {label}. Replace it with real content or TBD (owner: ...).")
        if not line.lstrip().startswith("|") and len(re.findall(r"\bTBD\b", text)) > len(TBD_OWNER.findall(text)):
            no_owner.append(ln)
        if TBD_PHRASE.search(text):
            phrases.append(ln)
        if BRACKET_PLACEHOLDER.search(text) and not CHECK.match(line):
            brackets.append(ln)
    if not blank:
        no_owner = sorted(no_owner + unowned_table_tbds(blocks))
    if placeholders:
        errors.append(f"{len(placeholders)} line(s) still contain <placeholders> (lines {numbers(placeholders)}). "
                      "Replace them with real content or TBD (owner: ...).")
    if no_owner:
        errors.append(f"TBD without an owner on line(s) {numbers(no_owner)}. Write TBD (owner: name or role).")
    if phrases:
        errors.append(f"Line(s) {numbers(phrases)} say 'tbd' or 'to be decided/confirmed'. "
                      "Write TBD (owner: name or role) so the gap is tracked.")
    if brackets:
        warnings.append(f"Line(s) {numbers(brackets)} contain [Bracketed text] that may be a placeholder.")

    ids = {}
    for b in blocks:
        candidates = []
        if b["kind"] == "heading" and b["level"] >= 3:
            candidates.append(b["text"])
        elif b["kind"] == "table":
            candidates += [row[0] for row in b["rows"][1:] if row]
        for c in candidates:
            m = ID_DEF.match(plain(c).strip())
            if m:
                ids.setdefault(m.group(1), []).append(b["line"])
    for id_, where in ids.items():
        if len(where) > 1:
            errors.append(f"{id_} is defined more than once (lines {', '.join(map(str, where))}). IDs must be unique.")

    if not blank:
        for b in blocks:
            if b["kind"] != "image":
                continue
            path = Path(b["path"])
            path = path if path.is_absolute() else src_dir / path
            if not path.is_file():
                errors.append(f"Line {b['line']}: image not found: {b['path']}")
            elif path.suffix.lower() not in IMAGE_TYPES:
                errors.append(f"Line {b['line']}: {path.suffix} images aren't supported; use PNG or JPG.")


def section_body(blocks, i):
    """Blocks under the heading at index I, up to the next heading of the same or a higher level."""
    level, out = blocks[i]["level"], []
    for b in blocks[i + 1:]:
        if b["kind"] == "heading" and b["level"] <= level:
            break
        out.append(b)
    return out


def check_content(blocks, doc_type, errors, warnings):
    for i, b in enumerate(blocks):
        if b["kind"] == "heading" and not any(x["kind"] not in ("heading", "pagebreak") for x in section_body(blocks, i)):
            errors.append(f"Line {b['line']}: '{plain(b['text'])}' has no content. Add content, TBD (owner: ...) or "
                          "'Not applicable: <reason>', or remove the sub-section if it is optional.")
        if b["kind"] in ("para", "bullet"):
            m = NOT_APPLICABLE.match(plain(b["text"]).strip())
            if m:
                reason = m.group(1).lstrip(":.-– ").strip()
                if len(reason.split()) < 3:
                    errors.append(f"Line {b['line']}: write 'Not applicable: <reason>' with a real reason.")
                elif UNDECIDED.search(reason):
                    warnings.append(f"Line {b['line']}: this sounds undecided rather than not applicable. "
                                    "Use TBD (owner: ...) for things that haven't been decided yet.")
    if doc_type == "SRS":
        check_requirements(blocks, errors, warnings)


def check_requirements(blocks, errors, warnings):
    vague = ("uses vague wording ('{}'). Give a measurable target from a source, or TBD (owner: ...) "
             "with an open question proposing a value.")
    for i, b in enumerate(blocks):
        if b["kind"] != "heading" or b["level"] != 3:
            continue
        m = FR_HEADING.match(plain(b["text"]).strip())
        if not m or "withdrawn" in b["text"].lower():
            continue
        body = section_body(blocks, i)
        fields = {}
        for j, x in enumerate(body):
            if x["kind"] == "bullet" and x["level"] == 0:
                key, _, rest = plain(x["text"]).strip().partition(":")
                fields[key.strip().lower()] = (rest.strip(), j)
        missing = []
        if not re.match(r"(Must|Should|Could)\b", fields.get("priority", ("",))[0]):
            missing.append("a Priority of Must, Should or Could")
        description = fields.get("description", ("",))[0]
        if "shall" not in description.lower():
            missing.append("a Description using 'shall'")
        criteria = fields.get("acceptance criteria")
        nested = criteria and any(y["kind"] in ("number", "bullet") and y.get("level", 0) >= 1
                                  for y in body[criteria[1] + 1:criteria[1] + 2])
        if not criteria or not (criteria[0] or nested):
            missing.append("acceptance criteria")
        if not fields.get("source", ("",))[0]:
            missing.append("a Source")
        if missing:
            errors.append(f"Line {b['line']}: {m.group(1)} needs {', '.join(missing)}.")
        word = VAGUE.search(description)
        if word:
            warnings.append(f"Line {b['line']}: {m.group(1)} " + vague.format(word.group(1)))
    for b in blocks:
        if b["kind"] != "table" or not b["rows"]:
            continue
        header = [plain(c).strip().lower() for c in b["rows"][0]]
        target = header.index("target or measure") if "target or measure" in header else None
        for row in b["rows"][1:]:
            rid = plain(row[0]).strip() if row else ""
            if not re.match(r"NFR-\d{3}$", rid):
                continue
            word = VAGUE.search(plain(" ".join(row[1:])))
            if word:
                warnings.append(f"Line {b['line']}: {rid} " + vague.format(word.group(1)))
            value = plain(row[target]) if target is not None and target < len(row) else ""
            unmeasured = not re.search(r"\d", value) and (VAGUE.search(value) or len(value.split()) < 4)
            if target is not None and unmeasured and "TBD" not in value and "(Proposed" not in value:
                warnings.append(f"Line {b['line']}: {rid} has no measurable target. Use a number or named standard "
                                "from a source, or TBD (owner: ...).")


def named(value):
    v = (value or "").strip()
    return bool(v) and not v.upper().startswith("TBD") and v.lower() not in NOT_A_NAME and not PLACEHOLDER.search(v)


def open_questions(blocks):
    found = []
    for b in blocks:
        if b["kind"] != "table" or not b["rows"]:
            continue
        header = [plain(c).strip().lower() for c in b["rows"][0]]
        if "status" not in header:
            continue
        k = header.index("status")
        for row in b["rows"][1:]:
            rid = plain(row[0]).strip() if row else ""
            if re.match(r"OQ-\d{3}$", rid) and k < len(row) and plain(row[k]).strip().lower() == "open":
                found.append(rid)
    return found


def check_approval(meta, blocks, tbds, proposed, errors, warnings):
    status = meta.get("status")
    if status == "Approved":
        if not named(meta.get("reviewed_by")):
            errors.append(f"An approved document needs the reviewer's name in reviewed_by (found '{meta.get('reviewed_by', '')}').")
        if not named(meta.get("approved_by")):
            errors.append(f"An approved document needs the approver's name in approved_by (found '{meta.get('approved_by', '')}').")
        if not named(meta.get("approval_ref")):
            errors.append("An approved document needs approval_ref: where the approval is recorded "
                          "(e-signature ID, signed copy, or email and date).")
        if tbds:
            errors.append(f"An approved document can't contain open TBDs ({len(tbds)} found).")
        if proposed:
            errors.append(f"An approved document can't contain Proposed items ({len(proposed)} found). "
                          "Confirm or remove each one first.")
        questions = open_questions(blocks)
        if questions:
            errors.append(f"An approved document can't have open questions: {', '.join(questions)}.")
    else:
        if named(meta.get("approved_by")):
            warnings.append(f"approved_by names '{meta['approved_by']}' but the status is {status}. On a new version, "
                            "set approved_by back to TBD until that version is approved.")
        if meta.get("approval_ref"):
            warnings.append("approval_ref is set but the document isn't approved; it belongs to an earlier version.")


def check_rhn(blocks, errors):
    section, figures, report_named = "", False, False
    for b in blocks:
        if b["kind"] == "heading" and b["level"] == 1:
            section = plain(b["text"]).lower()
            continue
        if not section.startswith("6."):
            continue
        if b["kind"] == "table" and b["rows"]:
            header = [plain(c).strip().lower() for c in b["rows"][0]]
            columns = [k for k, h in enumerate(header) if h in ("executed", "passed", "failed")]
            figures |= any(k < len(row) and re.search(r"\d", row[k]) for row in b["rows"][1:] for k in columns)
        elif b["kind"] in ("para", "bullet") and "report" in b["text"].lower():
            evidence = b["text"].split(":", 1)[-1].strip()
            report_named |= bool(evidence) and "TBD" not in evidence and not PLACEHOLDER.search(evidence)
    if figures and not report_named:
        errors.append("Test summary: test figures need a named test report (file name or link). "
                      "If there is no report, set the figures to TBD (owner: ...).")


def short(text, n=110):
    return text if len(text) <= n else text[: n - 3] + "..."


def report(src, meta, sections, errors, warnings, tbds, proposed, ticks, blank):
    print(f"{src.name}: {meta.get('type') or '?'} v{meta.get('version', '?')}, {meta.get('status', '?')}")
    if sections:
        print(f"  Structure: {sections[0]} of {sections[1]} standard sections present")
    for e in errors:
        print(f"  ERROR: {e}")
    for w in warnings:
        print(f"  WARNING: {w}")
    if blank:
        return
    if tbds:
        print(f"  {len(tbds)} TBD item(s):")
        for ln, text in tbds:
            print(f"    line {ln}: {short(text)}")
    if proposed:
        print(f"  {len(proposed)} Proposed item(s), which the user must confirm or reject:")
        for ln, text in proposed:
            print(f"    line {ln}: {short(text)}")
    if ticks:
        print(f"  {len(ticks)} ticked checklist item(s); ticks need evidence from the user:")
        for ln, text in ticks:
            print(f"    line {ln}: {short(text)}")


def secret_files_near(src):
    """Names of secret-bearing files in the source's folder (and, for a docs/ folder, the project folder above it)."""
    roots = [src.parent] + ([src.parent.parent] if src.parent.name.lower() == "docs" else [])
    found, seen = [], 0
    for root in roots:
        for path in root.rglob("*"):
            if seen > 5000 or len(found) >= 20:
                break
            seen += 1
            if SKIP_DIRS.intersection(path.relative_to(root).parts[:-1]) or not path.is_file():
                continue
            rel = path.relative_to(roots[-1]).as_posix()
            if SECRET_FILES.match(path.name) and rel not in found:
                found.append(rel)
    return found


def handover(meta, out, tbds, proposed, questions, ticks, secret_files):
    owners = {}
    for _, text in tbds:
        for owner in TBD_OWNER.findall(text):
            owners[owner.strip()] = owners.get(owner.strip(), 0) + 1
    print("HAND-OVER: your final message to the user must cover all of the following:")
    print(f"  - The file: {out.name} ({meta['type']} v{meta['version']}, {meta['status']})")
    if meta.get("project_code_status", "").lower() == "provisional":
        print(f"  - The project code {meta['project_code']} is provisional: it must be agreed before approval")
    if secret_files:
        print("  - Files that usually hold secrets are in the project folder (" + ", ".join(secret_files) + "). "
              "Tell the user they should not be shared or committed and may need rotating; do not open them.")
    if owners:
        print("  - Open TBD items by owner: " + "; ".join(f"{o} ({n})" for o, n in owners.items()))
    else:
        print("  - No open TBD items.")
    if proposed:
        print(f"  - {len(proposed)} Proposed item(s), highlighted in turquoise, that the user must confirm or reject")
    if questions:
        print(f"  - Open questions: {', '.join(questions)}")
    if ticks:
        print(f"  - {len(ticks)} ticked checklist item(s): ask the user to confirm the evidence for each")
    print("  - Word asks to update fields when the file is first opened; answering Yes fills in the table of contents.")


# ---------------------------------------------------------------- Word helpers

def make_el(tag, **attrs):
    el = OxmlElement(tag)
    for key, value in attrs.items():
        el.set(qn(f"w:{key}"), str(value))
    return el


def insert_ordered(parent, child, order):
    """Insert CHILD into PARENT at its schema position, replacing any element with the same tag."""
    for old in parent.findall(child.tag):
        parent.remove(old)
    tags = [qn(t) for t in order]
    for tag in tags[tags.index(child.tag) + 1:]:
        found = parent.find(tag)
        if found is not None:
            found.addprevious(child)
            return child
    parent.append(child)
    return child


def shade(cell, fill):
    insert_ordered(cell._tc.get_or_add_tcPr(), make_el("w:shd", val="clear", color="auto", fill=fill), TCPR_ORDER)


def fld_char(kind):
    return make_el("w:fldChar", fldCharType=kind)


class Ctx:
    def __init__(self, doc, brand, src_dir):
        self.doc, self.brand, self.src_dir = doc, brand, src_dir
        sec = doc.sections[0]
        self.text_width = sec.page_width - sec.left_margin - sec.right_margin
        self.break_next = False
        self.figures = 0

    def par(self, style, text=None):
        p = self.doc.add_paragraph(style=style)
        if self.break_next:
            p.paragraph_format.page_break_before = True
            self.break_next = False
        if text:
            add_inline(self, p, text)
        return p


def add_text(par, text, bold=False, italic=False, mono=False):
    for k, chunk in enumerate(BR.split(text)):
        if k:
            par.add_run().add_break()
        for piece in MARKERS.split(chunk):
            if not piece:
                continue
            run = par.add_run(piece)
            if bold:
                run.bold = True
            if italic:
                run.italic = True
            if mono:
                run.font.name = "Consolas"
            if TBD.fullmatch(piece):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            elif PROPOSED.fullmatch(piece):
                run.font.highlight_color = WD_COLOR_INDEX.TURQUOISE


def add_link(ctx, par, text, url):
    if url.startswith("#"):
        add_text(par, text)
        return
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), par.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run, rpr = OxmlElement("w:r"), OxmlElement("w:rPr")
    rpr.append(make_el("w:color", val=ctx.brand["accent_color"]))
    rpr.append(make_el("w:u", val="single"))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(rpr)
    run.append(t)
    link.append(run)
    par._p.append(link)


def add_inline(ctx, par, text, bold=False, italic=False):
    pos = 0
    for m in INLINE.finditer(text):
        add_text(par, text[pos:m.start()], bold, italic)
        if m.group("b") is not None:
            add_inline(ctx, par, m.group("b"), True, italic)
        elif m.group("c") is not None:
            add_text(par, m.group("c"), bold, italic, mono=True)
        elif m.group("lt") is not None:
            add_link(ctx, par, m.group("lt"), m.group("lu"))
        else:
            add_inline(ctx, par, m.group("i"), bold, True)
        pos = m.end()
    add_text(par, text[pos:], bold, italic)


def min_width(texts):
    """Roughly the width the longest unbreakable word needs at table text size, so IDs and dates don't wrap."""
    longest = 0
    for text in texts:
        t = plain(text).strip()
        longest = max([longest] + [len(p) for p in ([t] if DATE.fullmatch(t) else t.split())])
    return Cm(0.2 * min(longest, 30) + 0.45)


def fit_widths(total, weights, minimums):
    """Share TOTAL between columns by WEIGHTS, but give every column at least its minimum."""
    fixed = {}
    while True:
        free = [j for j in range(len(weights)) if j not in fixed]
        room = total - sum(fixed.values())
        share = sum(weights[j] for j in free)
        widths = {j: room * weights[j] / share for j in free}
        short = [j for j in free if widths[j] < minimums[j]]
        if not short or len(short) == len(free):
            break
        for j in short:
            fixed[j] = minimums[j]
    widths.update(fixed)
    return [Emu(int(widths[j])) for j in range(len(weights))]


def add_table(ctx, rows, header=True, weights=None, label_col=False):
    doc, brand = ctx.doc, ctx.brand
    ncols = max(len(r) for r in rows)
    rows = [list(r) + [""] * (ncols - len(r)) for r in rows]
    if ctx.break_next:
        ctx.par("Spacer")
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = doc.styles["Table Grid"]
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(make_el(f"w:{edge}", val="single", sz=4, space=0, color="BFBFBF"))
    insert_ordered(table._tbl.tblPr, borders, TBLPR_ORDER)
    if weights is None:
        weights = [min(max(len(plain(r[j])) for r in rows), 50) + 4 for j in range(ncols)]
    widths = fit_widths(ctx.text_width, weights, [min_width(r[j] for r in rows) for j in range(ncols)])
    for j, column in enumerate(table.columns):
        column.width = widths[j]
    for i, row in enumerate(rows):
        cells = table.rows[i].cells
        for j, text in enumerate(row):
            cell = cells[j]
            cell.width = widths[j]
            p = cell.paragraphs[0]
            p.style = doc.styles["Table Text"]
            add_inline(ctx, p, text)
            if header and i == 0:
                shade(cell, brand["primary_color"])
                for run in p.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif label_col and j == 0:
                shade(cell, brand["light_color"])
                for run in p.runs:
                    run.bold = True
    if header:
        table.rows[0]._tr.get_or_add_trPr().append(make_el("w:tblHeader", val="true"))
    doc.add_paragraph(style="Spacer")
    return table


def add_image(ctx, block):
    path = Path(block["path"])
    path = path if path.is_absolute() else ctx.src_dir / path
    p = ctx.par("Figure")
    if not path.is_file() or path.suffix.lower() not in IMAGE_TYPES:  # only reachable with --blank
        add_text(p, f"[Image: {block['alt']}]")
        return
    pic = p.add_run().add_picture(str(path))
    scale = min(1.0, ctx.text_width / pic.width, Cm(18) / pic.height)
    if scale < 1:
        pic.width, pic.height = Emu(int(pic.width * scale)), Emu(int(pic.height * scale))
    ctx.figures += 1
    ctx.par("Caption", f"Figure {ctx.figures}: {block['alt']}")


def render(ctx, blocks):
    for b in blocks:
        kind = b["kind"]
        if kind == "heading":
            ctx.par(f"Heading {b['level']}", b["text"])
        elif kind == "para":
            ctx.par("Normal", b["text"])
        elif kind == "bullet":
            ctx.par(("List Bullet", "List Bullet 2", "List Bullet 3")[b["level"]], b["text"])
        elif kind == "number":
            p = ctx.par(("Numbered Item", "Numbered Item 2", "Numbered Item 3")[b["level"]])
            p.add_run(f"{b['num']}.\t")
            add_inline(ctx, p, b["text"])
        elif kind == "check":
            p = ctx.par("Checklist")
            p.add_run(("☑" if b["done"] else "☐") + "\t")
            add_inline(ctx, p, b["text"])
        elif kind == "table":
            add_table(ctx, b["rows"])
        elif kind == "code":
            p = ctx.par("Code Block")
            for k, line in enumerate(b["lines"]):
                if k:
                    p.add_run().add_break()
                p.add_run(line)
        elif kind == "note":
            ctx.par("Note", b["text"])
        elif kind == "image":
            add_image(ctx, b)
        elif kind == "pagebreak":
            ctx.break_next = True


# ---------------------------------------------------------------- document parts

def prepare_meta(meta):
    m = dict(meta)
    label = DOC_TYPES[m["type"]][0]
    m["title"] = m.get("title") or label
    m["type_label"] = f"{label} ({m['type']})"
    for key in ("reviewed_by", "approved_by"):
        m[key] = m.get(key) or "TBD"
    m["classification"] = m.get("classification") or "Confidential"
    for key in REQUIRED:
        m.setdefault(key, f"<{key}>")
    return m


def fill_tokens(doc, tokens):
    for sec in doc.sections:
        for part in (sec.header, sec.footer):
            for p in part.paragraphs:
                for run in p.runs:
                    for key, value in tokens.items():
                        if key in run.text:
                            run.text = run.text.replace(key, value)


def add_cover(ctx, meta):
    brand = ctx.brand
    p = ctx.par("Cover Brand")
    logo = ASSETS / "logo.png"
    if logo.is_file():
        p.add_run().add_picture(str(logo), height=Cm(1.6))
    else:
        p.add_run(brand["company"].upper())
    if meta.get("client_logo"):
        # Client logo on the right of the same line; marked so docx_to_md.py can recover it.
        path = Path(meta["client_logo"])
        path = path if path.is_absolute() else ctx.src_dir / path
        p.paragraph_format.tab_stops.add_tab_stop(ctx.text_width, WD_TAB_ALIGNMENT.RIGHT)
        p.add_run("\t")
        pic = p.add_run().add_picture(str(path), height=Cm(1.6))
        if pic.width > Cm(6):
            pic.width, pic.height = Cm(6), Emu(int(pic.height * Cm(6) / pic.width))
        pic._inline.docPr.set("descr", "client-logo")
    ctx.par("Title", meta["title"])
    ctx.par("Subtitle", meta["project"])
    shown = dict(meta)
    if shown.get("project_code_status", "").lower() == "provisional":
        shown["project_code_status"] = "Provisional: to be agreed with the client"
    optional = ("approval_ref", "project_code_status")
    rows = [[label, shown.get(key, "")] for label, key in COVER_FIELDS
            if key not in optional or (meta.get(key) and not (key == "project_code_status" and meta[key].lower() == "agreed"))]
    add_table(ctx, rows, header=False, weights=[1, 2.6], label_col=True)
    notice = brand["confidentiality_notice"]
    for token, value in (("{classification}", meta["classification"]), ("{company}", brand["company"]),
                         ("{client}", meta["client"])):
        notice = notice.replace(token, value)
    ctx.par("Cover Notice", notice)


def add_revisions(ctx, revisions):
    ctx.break_next = True
    ctx.par("Front Heading", "Revision history")
    rows = [["Version", "Date", "Author", "Changes"]] + [fields[:4] for fields, _ in revisions]
    if len(rows) == 1:
        rows.append(["", "", "", ""])
    add_table(ctx, rows, weights=[1, 1.4, 1.9, 4.5])


def add_toc(ctx, headings):
    ctx.break_next = True
    ctx.par("Front Heading", "Contents")
    entries = [(h["level"], plain(h["text"])) for h in headings]
    if not entries:
        entries = [(1, "Right-click here and choose Update Field to build the table of contents.")]
    # The entries are a preview; Word replaces them, with page numbers, when fields update.
    paras = [ctx.par(f"toc {level}") for level, _ in entries]
    paras[0].add_run()._r.append(fld_char("begin"))
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    paras[0].add_run()._r.append(instr)
    paras[0].add_run()._r.append(fld_char("separate"))
    for p, (_, text) in zip(paras, entries):
        p.add_run(text)
    paras[-1].add_run()._r.append(fld_char("end"))


def build(meta, revisions, blocks, brand, src_dir, out_path, blank):
    doc = Document(str(ASSETS / "panasa-template.docx"))
    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    ctx = Ctx(doc, brand, src_dir)
    if blank:
        doc_ref = f"{meta['type']} template"
    else:
        provisional = " (provisional code)" if meta.get("project_code_status", "").lower() == "provisional" else ""
        doc_ref = f"{meta['project_code']}{provisional} · {meta['type']} · v{meta['version']} · {meta['status']}"
    fill_tokens(doc, {"{{DOC_REF}}": doc_ref, "{{CLASSIFICATION}}": meta["classification"]})

    add_cover(ctx, meta)
    add_revisions(ctx, revisions)
    add_toc(ctx, [b for b in blocks if b["kind"] == "heading" and b["level"] <= 2])
    ctx.break_next = True
    render(ctx, blocks)

    insert_ordered(doc.settings.element, make_el("w:updateFields", val="true"), SETTINGS_ORDER)
    now = datetime.now(timezone.utc).replace(tzinfo=None, microsecond=0)
    props = doc.core_properties
    props.title = f"{meta['title']} - {meta['project']}"
    props.subject = meta["project"]
    props.author = props.last_modified_by = meta["prepared_by"]
    props.keywords = f"{meta['project_code']}; {meta['type']}"
    props.category = meta["type_label"]
    props.content_status = meta["status"]
    props.version = meta["version"]
    props.comments = "Built with the Panasa project-documentation skill."
    props.created = props.modified = now
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    ap = argparse.ArgumentParser(description="Build a Panasa-branded Word document from a Markdown source.")
    ap.add_argument("source", help="Markdown source file")
    ap.add_argument("--out", help="output folder or .docx file (default: next to the source)")
    ap.add_argument("--check", action="store_true", help="validate only; write nothing")
    ap.add_argument("--blank", action="store_true", help="build a fill-in-by-hand template")
    ap.add_argument("--force", action="store_true", help="overwrite an existing approved .docx")
    args = ap.parse_args()

    src = Path(args.source).resolve()
    if not src.is_file():
        print(f"ERROR: {src} not found.")
        return 1
    text = src.read_text(encoding="utf-8-sig").replace("\r\n", "\n")
    meta, revisions, body_lines, first_line, errors = parse_source(text)
    meta["type"] = meta.get("type", "").upper()
    clean_lines = strip_comments("\n".join(body_lines)).split("\n")
    blocks = parse_blocks(clean_lines, first_line)
    warnings, tbds, proposed, sections = [], [], [], None
    check_secrets(text.split("\n"), errors)
    if meta["type"] not in DOC_TYPES:
        errors.append(f"type must be one of {', '.join(DOC_TYPES)} (found '{meta['type']}').")
    else:
        sections = check_structure(blocks, expected_sections(meta["type"]), errors, warnings)
        if not args.blank:
            check_meta(meta, revisions, src.parent, errors, warnings)
        check_body(blocks, clean_lines, first_line, src.parent, args.blank, errors, warnings, tbds, proposed)
        if not args.blank:
            check_content(blocks, meta["type"], errors, warnings)
            if meta["type"] == "RHN":
                check_rhn(blocks, errors)
            check_approval(meta, blocks, tbds, proposed, errors, warnings)
    ticks = [(b["line"], plain(b["text"])) for b in blocks if b["kind"] == "check" and b["done"]]
    report(src, meta, sections, errors, warnings, tbds, proposed, ticks, args.blank)
    if errors:
        print("Not built: fix the errors above and run again.")
        return 1
    if args.check:
        return 0

    brand = json.loads((ASSETS / "brand.json").read_text(encoding="utf-8"))
    meta = prepare_meta(meta)
    if args.blank:
        name = f"Panasa_{meta['type']}_template.docx"
    else:
        name = f"{meta['project_code']}_{meta['type']}_v{meta['version']}.docx"
    out = Path(args.out).resolve() if args.out else src.parent
    if out.suffix.lower() != ".docx":
        out = out / name
    if out.exists() and not args.force:
        try:
            previous = Document(str(out)).core_properties.content_status
        except Exception:
            previous = ""
        if previous == "Approved":
            print(f"ERROR: {out.name} is an approved document and is never overwritten. Raise the version for "
                  "any change, or use --force only if this file was never issued.")
            return 1
    try:
        build(meta, revisions, blocks, brand, src.parent, out, args.blank)
    except PermissionError:
        print(f"ERROR: can't write {out}. If it is open in Word, close it and run again.")
        return 1
    print(f"Built: {out}")
    if not args.blank:
        handover(meta, out, tbds, proposed, open_questions(blocks), ticks, secret_files_near(src))
    return 0


if __name__ == "__main__":
    sys.exit(main())
